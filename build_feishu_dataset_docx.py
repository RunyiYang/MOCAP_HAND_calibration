#!/usr/bin/env python3
"""Build an image-embedded DOCX version of the Feishu dataset note."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "outputs" / "feishu_dataset_understanding"
ASSETS = REPORT / "assets"
OUTPUT = REPORT / "Dataset理解_飞书版.docx"

INK = "172033"
BLUE = "426FE3"
GREEN = "29A76D"
YELLOW = "E4AE23"
LIGHT_BLUE = "EEF4FF"
LIGHT_GREEN = "EDFFF5"
LIGHT_YELLOW = "FFF7DF"
LIGHT_GRAY = "F5F7FB"
LINE = "DDE4EE"


def set_east_asia_font(run, name: str = "Noto Sans CJK SC") -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 130, start: int = 150, bottom: int = 130, end: int = 150) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run, *, bold: bool = False, size: float = 10.5, color: str = INK) -> None:
    set_east_asia_font(run)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_body(document: Document, text: str = "", *, bold_prefix: str | None = None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.35
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        style_run(first, bold=True)
        second = paragraph.add_run(text[len(bold_prefix) :])
        style_run(second)
    else:
        run = paragraph.add_run(text)
        style_run(run)
    return paragraph


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.25
    style_run(paragraph.add_run(text))


def add_number(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    style_run(paragraph.add_run(text))


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    style_run(run, bold=True, size=20 if level == 1 else 14, color=INK if level == 1 else BLUE)


def add_picture(document: Document, filename: str, caption: str, width_cm: float = 17.2) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run().add_picture(str(ASSETS / filename), width=Cm(width_cm))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    style_run(cap.add_run(caption), size=8.5, color="647087")


def add_callout(document: Document, title: str, body: str, *, fill: str, accent: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell, accent, "10")
    set_cell_margins(cell, 180, 210, 180, 210)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    style_run(p.add_run(title), bold=True, size=12, color=INK)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.25
    style_run(p2.add_run(body), size=10)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_data_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header_cells = table.rows[0].cells
    for index, label in enumerate(headers):
        cell = header_cells[index]
        shade_cell(cell, LIGHT_BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[index])
        paragraph = cell.paragraphs[0]
        style_run(paragraph.add_run(label), bold=True, size=9.2)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            if row_index % 2:
                shade_cell(cell, "FAFBFD")
            set_cell_border(cell)
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[index])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            style_run(paragraph.add_run(value), size=8.8)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.55)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Noto Sans CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    document.core_properties.title = "手套姿态 × MOCAP：Dataset 理解与首版视频可视化"
    document.core_properties.subject = "GT_calib dataset understanding"
    document.core_properties.author = "GT_calib review"

    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(eyebrow.add_run("GT_CALIB · DATASET NOTE · 2026-08-30"), bold=True, size=9, color=BLUE)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    style_run(title.add_run("手套姿态 × MOCAP\nDataset 理解与首版视频可视化"), bold=True, size=26, color=INK)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    style_run(
        subtitle.add_run("01、02、03 已完成时间同步、动捕世界系到 RGB 相机投影，以及双手黄绿骨架叠加。"),
        size=11.5,
        color="47546B",
    )
    add_picture(document, "overlay_triptych.png", "01 / 02 / 03 同一代表时刻：黄色 MOCAP 21 点，绿色 GLOVE 20 点。")
    add_callout(
        document,
        "页面结论",
        "黄色 MOCAP 21 点已有 world→RGB 投影链；绿色 GLOVE 20 点没有全局 wrist 平移，当前借用同步 MOCAP wrist，并使用每段固定旋转和尺度。因此黄绿距离是固定配准后的一致性差异，不是手套独立绝对精度。",
        fill=LIGHT_YELLOW,
        accent=YELLOW,
    )

    document.add_page_break()
    add_heading(document, "1. 这里的 MOCAP 到底是什么？")
    add_body(
        document,
        "MOCAP = motion capture（运动捕捉）数据本身。黄色骨架是我们对 MOCAP 输出做的可视化；它不是从 RGB 视频跑视觉模型得到，也不是动捕相机的原始画面。",
        bold_prefix="MOCAP = motion capture（运动捕捉）数据本身。",
    )
    add_picture(document, "mocap_pipeline.png", "上方是光学动捕链，下方是手套 IMU/运动学链；两者最终投到同一 RGB 像素坐标。")
    add_callout(
        document,
        "Human.cma",
        "当前黄色完整骨架来源。左右各 21 点：hand root + 五根手指 × 4 个节点，单位 mm，坐标在 mocap world。",
        fill=LIGHT_YELLOW,
        accent=YELLOW,
    )
    add_callout(
        document,
        "Body.cma",
        "只包含 LeftHand、RightHand、Forearm、Arm、Torso 等刚体中心，没有完整手指。Body 手刚体中心与 Human hand root 不是同一点，稳定相差约 69.5 mm。",
        fill=LIGHT_GRAY,
        accent="8D99AD",
    )
    add_body(document, "黄色 21 点也不等于 21 个实体 Marker 的逐点直接测量：光学相机观察 Marker，CMAvatar 结合骨骼模型求解连续 Skeleton。它是当前最强参考，但仍应称为 mocap reference，而不是未经限定的像素级绝对 Ground Truth。")

    add_heading(document, "2. Dataset 结构与每类文件的角色")
    add_picture(document, "dataset_map.png", "00 用于空间标定；01/02/03 是正式双手手套 × 动捕采集。")
    add_data_table(
        document,
        ["目录", "内容", "本次用途"],
        [
            ["视频/RGB.mp4", "1920×1080，固定 30 FPS", "最终叠加背景"],
            ["原始BAG与内参/*.bag", "原始 RGB-D + 设备时间", "真实 RGB timestamp；毫米深度"],
            ["同步校验/", "相机、Thor、CMAvatar 时钟", "同步三条数据流"],
            ["手套解算/*keypoints.csv", "20 点/手，单位 m", "绿色手形与 wrist orientation"],
            ["动捕/*Human.cma", "21 点/手，单位 mm", "黄色完整手骨架"],
            ["camera_to_world.json", "世界系/相机矩阵 + 内参", "mocap world 投到 RGB"],
        ],
        [4.7, 5.4, 6.4],
    )
    add_callout(document, "文件边界", "Depth.mp4 只是 640×576 可视化预览，毫米深度必须读原始 BAG；camera_aligned_pose 的 camera aligned 只表示时间对齐，文件内没有相机空间 xyz。", fill=LIGHT_BLUE, accent=BLUE)

    add_heading(document, "3. 三段正式数据与时间同步")
    add_data_table(
        document,
        ["段", "RGB", "Human", "Glove L/R", "同步 RGB", "Camera↔Mocap median/P95"],
        [
            ["01 / Take_000", "1815", "7215", "6943 / 6943", "1756", "2.087 / 4.081 ms"],
            ["02 / Take_001", "1812", "7482", "6791 / 6793", "1804", "2.118 / 3.932 ms"],
            ["03 / Take_002", "1810", "7349", "6688 / 6688", "1803", "2.117 / 3.952 ms"],
        ],
        [3.1, 1.6, 1.8, 2.6, 2.2, 5.2],
    )
    add_body(document, "三段共同区间均通过，camera→mocap 最大误差都小于 5 ms。精确帧链为：")
    formula = document.add_paragraph()
    formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    formula.paragraph_format.space_after = Pt(8)
    style_run(formula.add_run("MP4 frame i = BAG color message i\nBAG time → Thor monotonic → glove\nBAG time → CMAvatar time → Human.cma"), bold=True, size=10, color=BLUE)
    add_body(document, "MP4 标称 30 FPS，但 BAG RGB 实际约 29.912 Hz。若用 frame_index / 30 作为绝对时间，段末会累计约 176–178 ms 漂移；当前实现读取 BAG index 的真实时间。")

    add_heading(document, "4. 新增空间标定做了什么？")
    add_body(document, "00 段的 CS-400 标尺和 109 帧原始毫米深度，给出了右手世界系：原点位于长短臂虚拟交点正下方桌面，+X 沿长臂、+Y 沿短臂、+Z 向上。")
    add_data_table(
        document,
        ["证据", "数值", "正确解释"],
        [
            ["桌面拟合残差 P50", "0.767 mm", "桌面/深度几何质量"],
            ["桌面拟合残差 P95", "3.912 mm", "不是手关节误差"],
            ["长短臂实测夹角", "90.949°", "轴正交质量"],
            ["轴线最大 RGB 垂距", "约 2.43 px", "标记/轴线一致性"],
        ],
        [5.2, 3.0, 8.0],
    )
    add_picture(document, "reference_rgb_markers_and_world_axes.png", "00 标定参考：CS-400 标记、桌面世界轴和 45 mm 标记中心高度参考。")
    add_picture(document, "calibration_validation_triptych.png", "01/02/03 单帧验证：圆圈来自 Body.cma 左右手刚体中心。")
    add_callout(document, "标定证据边界", "三张 validation 每段只检查一个中间时刻的两个 Body 刚体中心，没有完整 Human 手关节的独立 2D 标注。所有 Human 投影点在画面内是 sanity check，不是像素精度证明。", fill=LIGHT_YELLOW, accent=YELLOW)

    add_heading(document, "5. 手套如何放到同一个世界系？")
    formula = document.add_paragraph()
    formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(formula.add_run("mocap_MCP − mocap_wrist ≈ scale · R · glove_world_MCP\nglove_in_world = mocap_wrist + scale · R · glove_world"), bold=True, size=10.5, color=BLUE)
    add_bullet(document, "每段、每只手只估计一个固定 R + scale；不会逐帧重新旋转去吸收差异。")
    add_bullet(document, "拟合只用 index/middle/ring/pinky 四个 MCP；四个非拇指指尖是 holdout。")
    add_bullet(document, "拇指为 20/21 点拓扑，绘制但不进入差异指标。")
    add_bullet(document, "每帧复制同步 mocap wrist 平移，因此不能评价手套的独立全局位置。")
    add_body(document, "图例：黄色空心 = MOCAP 21 点；绿色实心 = GLOVE 20 点；蓝线 = 非拇指指尖差异。")

    add_heading(document, "6. 首轮可视化看到了什么？")
    add_picture(document, "metrics_chart.png", "固定配准后的 3D 一致性差异 median；不是手套独立绝对精度。")
    add_body(document, "01 与 02 相近，03 明显更差。03 的非拇指关节 median 约 42–44 mm，指尖约 69–71 mm；视频中部分动作的绿色骨架横向分叉也更明显。", bold_prefix="01 与 02 相近，03 明显更差。")
    add_picture(document, "review_grid.png", "三段各 6 帧总览；从上到下依次为 01、02、03。")

    add_heading(document, "7. 当前结论与下一步")
    add_callout(
        document,
        "可以确认",
        "三条时钟链已贯通；新增外参方向、单位和公式自洽；黄色是 MOCAP 3D Skeleton 的 RGB 投影；03 的固定配准后一致性明显弱于 01/02。",
        fill=LIGHT_GREEN,
        accent=GREEN,
    )
    add_callout(
        document,
        "暂时不能宣称",
        "黄色逐关节是像素级绝对 GT；黄绿距离是手套独立绝对误差；手套自身提供相机中的全局 6DoF；相机移动后仍可复用当前固定外参。",
        fill=LIGHT_YELLOW,
        accent=YELLOW,
    )
    add_number(document, "审阅三段 H.264，优先标记 03 中明显分叉的时间区间。")
    add_number(document, "增加两种视图：保留 wrist orientation；共享 wrist orientation 只比较手指 articulation。")
    add_number(document, "如需绝对准确度，引入独立 2D 标注、相机刚体或其他未参与拟合的空间真值。")
    add_body(document, "验证：三段 1280×720 H.264 完整解码通过；RGB 帧数 1815 / 1812 / 1810；4 项数据契约与端到端测试通过。00 仅作为空间标定来源。")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
